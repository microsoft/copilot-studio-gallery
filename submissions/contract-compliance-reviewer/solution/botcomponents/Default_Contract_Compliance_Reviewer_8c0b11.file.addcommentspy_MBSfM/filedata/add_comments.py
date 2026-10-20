"""Add highlights + Word margin comments to a copy of a contract.

    python3 add_comments.py <original.docx> <issues.json> <new.docx>

issues.json is a list of {"text": "<exact text from one paragraph>", "comment": "<annotation>"}.
The comment prefix sets the text highlight colour and the comment author, so Word also
shows each category with its own balloon colour and reviewer name:
  LEGAL CONFLICT      -> red highlight,    author "Legal Compliance Review"
  ORG POLICY CONFLICT -> green highlight,  author "Org Policy Review"
  anything else       -> yellow highlight, author "Internal Consistency Review"
"""
import copy, json, sys
from datetime import datetime, timezone
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.text.run import Run

CATEGORIES = {  # prefix: (highlight, author, initials)
    "LEGAL CONFLICT": (WD_COLOR_INDEX.RED, "Legal Compliance Review", "LC"),
    "ORG POLICY CONFLICT": (WD_COLOR_INDEX.BRIGHT_GREEN, "Org Policy Review", "OP"),
}
DEFAULT = (WD_COLOR_INDEX.YELLOW, "Internal Consistency Review", "IC")


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def comments_part(doc):
    try:
        part = doc.part.part_related_by(RT.COMMENTS)
        if not hasattr(part, "element"):  # older python-docx loads existing comments as a plain part
            part._element, part.__class__ = parse_xml(part.blob), XmlPart
        root = part.element
    except KeyError:
        root = parse_xml('<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        doc.part.relate_to(XmlPart(PackURI("/word/comments.xml"), CT.WML_COMMENTS, root, doc.part.package), RT.COMMENTS)
    return root


def isolate(p, text):
    """Split runs so `text` is covered by whole runs (keeping their formatting); return those runs."""
    runs = lambda: [r for r in p.iter(qn("w:r")) if r.find(qn("w:t")) is not None]
    full = "".join(r.find(qn("w:t")).text or "" for r in runs())
    start = full.lower().find(text.lower())
    if start == -1:
        return None
    for cut in (start, start + len(text)):
        pos = 0
        for r in runs():
            t = r.find(qn("w:t"))
            s = t.text or ""
            if pos < cut < pos + len(s):
                right = copy.deepcopy(r)
                t.text, right.find(qn("w:t")).text = s[: cut - pos], s[cut - pos:]
                for x in (t, right.find(qn("w:t"))):
                    x.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.addnext(right)
                break
            pos += len(s)
    picked, pos = [], 0
    for r in runs():
        n = len(r.find(qn("w:t")).text or "")
        if pos < start + len(text) and pos + n > start and n:
            picked.append(r)
        pos += n
    return picked


def main(src, issues_file, out):
    if src == out:
        sys.exit("Output must be a new file - never overwrite the original.")
    doc = Document(src)
    root = comments_part(doc)
    cid = 1 + max([int(c.get(qn("w:id"))) for c in root.findall(qn("w:comment"))] or [-1])
    date = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    missing = []
    for issue in json.load(open(issues_file, encoding="utf-8-sig")):
        note = issue["comment"]
        runs = None
        for p in doc.element.body.iter(qn("w:p")):
            runs = isolate(p, issue["text"])
            if runs:
                break
        if not runs:  # text not found: keep the comment, attached to a note at the end
            missing.append(issue["text"])
            runs = [doc.add_paragraph().add_run("[Unlocated issue] " + issue["text"])._r]
        color, author, initials = next(
            (v for k, v in CATEGORIES.items() if note.lstrip("[ ").upper().startswith(k)), DEFAULT)
        for r in runs:
            Run(r, None).font.highlight_color = color
        c = el("w:comment", id=cid, author=author, initials=initials, date=date)
        cp = el("w:p")
        r1, r2, t = el("w:r"), el("w:r"), el("w:t")
        r1.append(el("w:annotationRef"))
        t.text = note
        r2.append(t)
        cp.extend([r1, r2])
        c.append(cp)
        root.append(c)
        runs[0].addprevious(el("w:commentRangeStart", id=cid))
        end = el("w:commentRangeEnd", id=cid)
        runs[-1].addnext(end)
        ref = el("w:r")
        ref.append(el("w:commentReference", id=cid))
        end.addnext(ref)
        cid += 1
    doc.save(out)
    print("Saved", out, "| unlocated:", missing or "none")


if __name__ == "__main__":
    main(*sys.argv[1:4])
